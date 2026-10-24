"""One owner-scoped timeline across stock, AI, billing, storage and account changes."""
import csv
import io
import json
import time
from datetime import date,datetime,timedelta,time as clock
from decimal import Decimal
from sqlalchemy import select,union_all,literal,String,cast,func,or_,case
from data import db_session
from data.activity import ActivityEvent
from data.product import ProductMutation,ProductScanUsage
from data.scan_credits import ScanRequest,ScanCreditPack
from data.billing import BillingOrder
from services import product_history as stock_history
from services import product_service as stock
from services.billing_service import PLANS

CATEGORIES={'recipes':'Рецепты','products':'Продукты','ai':'Распознавания и токены','payments':'Оплата и пакеты','subscription':'Подписка','storage':'Места хранения','account':'Аккаунт и настройки'}
ACTIONS={**stock_history.ACTIONS,'recipe_favorite':'Избранное','recipe_note':'Заметка к рецепту','recipe_steps':'Шаги готовки','recipe_nutrition':'Расчёт КБЖУ','recipe_nutrition_reset':'Сброс расчёта','recipe_steps_reset':'Сброс шагов','recipe_generated':'Создание рецепта с ИИ','recognition':'Распознавание','ocr_read':'Чтение этикетки OCR','barcode_lookup':'Поиск по штрихкоду','history_exported':'Экспорт истории','import_preview':'Предпросмотр файла','import_canceled':'Отмена импорта','payment':'Оплата',
 'storage_created':'Создание места','storage_changed':'Изменение места','storage_deleted':'Удаление места',
 'subscription_created':'Подписка создана','subscription_changed':'Изменение подписки','subscription_deleted':'Подписка удалена',
 'default_storage_created':'Место по умолчанию','default_storage_changed':'Место по умолчанию','default_storage_deleted':'Сброс места по умолчанию',
 'account_created':'Регистрация','account_changed':'Изменение аккаунта','account_deleted':'Удаление аккаунта','login':'Вход в аккаунт','logout':'Выход из аккаунта'}
STATUSES={'used':'Успешно','failed':'Ошибка','pending':'В обработке','interrupted':'Прервано','succeeded':'Успешно','canceled':'Отменено','rejected':'Отклонено','ready':'Не оплачено','created':'Создано','waiting_for_capture':'Ожидает подтверждения'}
KINDS={'recipe':'Рецепт с ИИ','label':'Этикетка по фото','product':'Продукт по внешнему виду','ocr':'Чтение этикетки OCR','barcode':'Поиск по штрихкоду'}
FIELD_LABELS={'name':'Название','kind':'Тип','location':'Расположение','note':'Заметка','plan':'Тариф','interval':'Период оплаты','extra_seats':'Дополнительные места','status':'Статус','period_start':'Начало периода','period_end':'Доступ до','cycle_amount':'Следующее списание, ₽','auto_renew':'Автопродление','card_last4':'Карта','default_storage_id':'Место по умолчанию (ID)','email':'Почта','is_verified':'Почта подтверждена','password_changed':'Пароль изменён','filename':'Файл','count':'Количество продуктов','format':'Формат экспорта','events':'Событий в экспорте'}


def text_fields(column,keys):
    result=literal('')
    for key in keys:result=result+literal(' ')+func.coalesce(column[key].as_string(),'')
    return result


def feed(owner):
    p=ProductMutation;a=ActivityEvent;b=BillingOrder;s=ScanRequest;old=ProductScanUsage
    def project(source,id,category,action,created,search,condition,status=None):
        return select(literal(source).label('source'),cast(id,String).label('id'),category.label('category'),action.label('action'),created.label('created_at'),search.label('search'),(status if status is not None else literal('')).label('status')).where(*condition)
    pp=project('product',p.id,literal('products'),p.action,p.created_at,
        text_fields(p.before,['name','brand','manufacturer','storage_name','note'])+text_fields(p.after,['name','brand','manufacturer','storage_name','operation_note','import_filename'])+case(stock_history.ACTIONS,value=p.action,else_=''),[p.owner_id==owner])
    aa=project('activity',a.id,a.category,a.action,a.created_at,
        text_fields(a.before,['name','location','note','email','filename'])+text_fields(a.after,['name','location','note','email','product_name','error','filename'])+case(ACTIONS,value=a.action,else_='')+case({k:v['name'] for k,v in PLANS.items()},value=a.after['plan'].as_string(),else_=''),[a.owner_id==owner],case((a.category=='ai',a.after['status'].as_string()),else_=''))
    bb=project('billing',b.id,literal('payments'),literal('payment'),func.coalesce(b.finished_at,b.created_at),
        cast(b.amount,String)+literal(' ')+b.plan+literal(' ')+b.kind+literal(' ')+case({k:v['name'] for k,v in PLANS.items()},value=b.plan,else_=''),[b.user_id==owner],b.status)
    ss=project('scan',s.id,literal('ai'),literal('recognition'),s.created_at,s.kind,
        [s.user_id==owner,~select(a.id).where(a.id==s.id,a.owner_id==owner).exists()],s.status)
    oo=project('legacy',old.id,literal('ai'),literal('recognition'),old.created_at,literal('Распознавание'),[old.user_id==owner],old.status)
    return union_all(pp,aa,bb,ss,oo).subquery()


def filtered(args,owner):
    values={k:str(args.get(k,'')).strip() for k in ['q','category','action','status','date_from','date_to','group_id','origin','storage_id']}
    values['q']=stock.clean(values['q'])[:120]
    if values['category'] and values['category'] not in CATEGORIES:raise ValueError('Выберите раздел из списка.')
    if values['action'] and values['action'] not in ACTIONS:raise ValueError('Выберите действие из списка.')
    import re
    for key in ['group_id','storage_id']:
        if values[key] and not re.fullmatch(r'[0-9]{1,18}',values[key]):raise ValueError('Неверный фильтр продукта или места.')
    if values['origin'] and not re.fullmatch(r'[0-9]{1,12}(?:\.[0-9]{1,9})?',values['origin']):raise ValueError('Неверный фильтр продукта.')
    if values['status'] and values['status'] not in STATUSES:raise ValueError('Выберите статус из списка.')
    dates={}
    for key in ['date_from','date_to']:
        if values[key]:
            try:
                dates[key]=date.fromisoformat(values[key])
                if not 1900<=dates[key].year<=2200:raise ValueError
            except ValueError:raise ValueError('Проверьте даты периода.')
    if dates.get('date_from') and dates.get('date_to') and dates['date_from']>dates['date_to']:raise ValueError('Начало периода позже конца.')
    source=feed(owner);conditions=[]
    if values['category']:conditions.append(source.c.category==values['category'])
    if values['action']:conditions.append(source.c.action==values['action'])
    effective_status=case(((source.c.status=='pending') & (source.c.category=='ai') & (source.c.created_at<time.time()-600),'interrupted'),else_=source.c.status)
    if values['status']:conditions.append(effective_status==values['status'])
    if values['q']:conditions.append(func.unicode_casefold(source.c.search).contains(values['q'].casefold(),autoescape=True))
    for key,day in dates.items():
        start=datetime.combine(day,clock.min,tzinfo=stock_history.MOSCOW)
        conditions.append(source.c.created_at>=start.timestamp() if key=='date_from' else source.c.created_at<(start+timedelta(days=1)).timestamp())
    if values['group_id'] or values['storage_id']:
        matched=select(ProductMutation.id).where(ProductMutation.owner_id==owner)
        if values['group_id']:
            parts=[]
            for column in [ProductMutation.before,ProductMutation.after]:
                part=column['group_id'].as_integer()==int(values['group_id'])
                if values['origin']:part=part & (column['group_created_at'].as_float()==float(values['origin']))
                parts.append(part)
            matched=matched.where(or_(*parts))
        if values['storage_id']:matched=matched.where(or_(*[c['storage_id'].as_integer()==int(values['storage_id']) for c in [ProductMutation.before,ProductMutation.after]]))
        conditions.extend([source.c.source=='product',source.c.id.in_(select(cast(matched.subquery().c.id,String)))])
    return values,source,conditions


def money(value):return format(Decimal(value)/100,'.2f').replace('.',',')+' ₽' if value is not None else '—'


def token_usage(usage):
    calls=usage.get('calls',[]) if isinstance(usage,dict) else []
    totals=[]
    for call in calls:
        tokens=call.get('tokens',{})
        value=tokens.get('total_tokens')
        if value is None and 'prompt_tokens' in tokens and 'completion_tokens' in tokens:value=tokens['prompt_tokens']+tokens['completion_tokens']
        totals.append(value)
    # A timeout can consume tokens at the provider without producing a usage record.
    complete=bool(calls) and not usage.get('unreported_attempts') and all(v is not None for v in totals)
    known=sum(v for v in totals if v is not None)
    return dict(total=known if complete else None,known=known,calls=len(calls),partial=not complete and bool(known))


def settings_value(key,value):
    if value is None or value=='':return '—'
    if isinstance(value,bool):return 'Да' if value else 'Нет'
    if key=='plan':return PLANS.get(value,{}).get('name',value)
    if key=='interval':return {'month':'Месяц','year':'Год'}.get(value,value)
    if key=='status':return {'active':'Активна','trialing':'Пробный период','inactive':'Не активна','past_due':'Нужна оплата'}.get(value,value)
    if key=='kind':return {'fridge':'Холодильник','freezer':'Морозильник','pantry':'Шкаф','other':'Другое'}.get(value,value)
    if key in ['period_start','period_end']:return datetime.fromtimestamp(value,stock_history.MOSCOW).strftime('%d.%m.%Y %H:%M МСК')
    if key=='cycle_amount':return money(value)
    if key=='card_last4':return '•••• '+str(value)
    return str(value)


def hydrate(ses,items,owner):
    by_source={k:[] for k in ['product','activity','billing','scan','legacy']}
    for item in items:by_source[item.source].append(int(item.id) if item.source=='product' else item.id)
    models={'product':(ProductMutation,'owner_id'),'activity':(ActivityEvent,'owner_id'),'billing':(BillingOrder,'user_id'),'scan':(ScanRequest,'user_id'),'legacy':(ProductScanUsage,'user_id')}
    objects={}
    for source,ids in by_source.items():
        if ids:
            cls,owner_column=models[source]
            # Split to keep within SQLite parameter limits, including older builds.
            for offset in range(0,len(ids),400):
                for obj in ses.scalars(select(cls).where(cls.id.in_(ids[offset:offset+400]),getattr(cls,owner_column)==owner)):
                    objects[(source,str(obj.id))]=obj
    pack_ids=[o.id for (s,_),o in objects.items() if s=='billing' and o.kind=='scan_pack']
    packs={}
    for offset in range(0,len(pack_ids),400):
        packs.update({p.order_id:p for p in ses.scalars(select(ScanCreditPack).where(ScanCreditPack.user_id==owner,ScanCreditPack.order_id.in_(pack_ids[offset:offset+400])))})
    results=[]
    for item in items:
        obj=objects.get((item.source,str(item.id)))
        if obj is None:continue
        stamp=item.created_at
        v=dict(key=item.source+'-'+str(item.id),source=item.source,id=str(item.id),category=item.category,category_label=CATEGORIES[item.category],action=item.action,
            title=ACTIONS.get(item.action,item.action),created_at=stamp,timestamp=datetime.fromtimestamp(stamp,stock_history.MOSCOW).strftime('%d.%m.%Y %H:%M:%S МСК'),
            day=datetime.fromtimestamp(stamp,stock_history.MOSCOW).strftime('%d.%m.%Y'),name='',summary='',status='',status_label='',details=[],changes=[],before={},after={},tokens=None)
        if item.source=='product':
            stockview=stock_history.view(obj);v.update(name=stockview['name'],changes=stockview['changes'],before=obj.before,after=obj.after,stock=stockview)
            v['summary']=stockview['quantity']+' · '+stockview['place']
            if stockview['destination'] and stockview['destination']!=stockview['place']:v['summary']+=' → '+stockview['destination']
            if stockview['reason']:v['details'].append(('Причина',stockview['reason']))
            if stockview['note']:v['details'].append(('Комментарий',stockview['note']))
            if stockview['filename']:v['details'].append(('Файл',stockview['filename']))
        elif item.source=='activity' and obj.category!='ai':
            v.update(before=obj.before,after=obj.after,name=obj.after.get('name') or obj.before.get('name',''))
            for key in FIELD_LABELS:
                if obj.before.get(key)!=obj.after.get(key):v['changes'].append(dict(label=FIELD_LABELS[key],before=settings_value(key,obj.before.get(key)),after=settings_value(key,obj.after.get(key))))
            v['summary']=''
        elif item.source=='billing':
            pack=packs.get(obj.id);v['status']=obj.status
            v['name']='Пакет распознаваний' if obj.kind=='scan_pack' else PLANS.get(obj.plan,{}).get('name','Подписка')
            v['title']={'scan_pack':'Покупка пакета','bind':'Подтверждение карты','renewal':'Продление подписки','purchase':'Оформление подписки'}.get(obj.kind,'Оплата')
            v['summary']=money(obj.amount)
            v['details']=[('Номер операции',obj.id),('Статус',STATUSES.get(obj.status,obj.status))]
            if pack:v['details'].extend([('Тип пакета',KINDS.get(pack.kind,pack.kind)),('Распознаваний',str(pack.quantity)),('Начислено',str(pack.quantity) if obj.status=='succeeded' else '0')])
            elif obj.kind!='scan_pack':v['details'].extend([('Период','Год' if obj.interval=='year' else 'Месяц'),('Дополнительных мест',str(obj.extra_seats))])
            if obj.card_last4:v['details'].append(('Карта','•••• '+obj.card_last4))
            if obj.refund_id:v['details'].append(('Возврат проверки карты','Оформлен'))
            v['after']=dict(kind=obj.kind,plan=obj.plan,interval=obj.interval,status=obj.status,amount_kopecks=obj.amount,extra_seats=obj.extra_seats,credited=pack.quantity if pack and obj.status=='succeeded' else 0)
        else:
            if item.source=='activity':a=obj.after;kind=a.get('kind');status=a.get('status','pending');usage=a.get('usage',{});charged=a.get('charged',0)
            else:kind=obj.kind if item.source=='scan' else None;status=obj.status;usage={};charged=1 if status=='used' else 0;a=dict(kind=kind,status=status,charged=charged,credit_pack_id=getattr(obj,'credit_pack_id',None))
            if status=='pending' and stamp<time.time()-600:status='interrupted'
            v['status']=status;v['name']=a.get('product_name','');v['title']=KINDS.get(kind,'Распознавание · старый формат')
            v['tokens']=token_usage(usage);v['after']=a
            v['summary']=f'Списано распознаваний: {charged}'
            v['details']=[('Источник лимита','Докупленный пакет' if a.get('credit_pack_id') else 'Включено в подписку'),('Списано распознаваний',str(charged)),('Режим',KINDS.get(kind,'Не сохранён в старой версии'))]
            t=v['tokens'];v['details'].append(('Токены GigaChat',str(t['total']) if t['total'] is not None else f"Не менее {t['known']} (неполные данные)" if t['partial'] else 'Нет данных о фактическом расходе'))
            if t['calls']:v['details'].append(('Получено ответов модели',str(t['calls'])))
            for i,call in enumerate(usage.get('calls',[]),1):
                if call.get('model'):v['details'].append((f'Ответ {i} · модель',call['model']))
                for key,label in [('prompt_tokens','входные токены'),('completion_tokens','токены ответа'),('total_tokens','всего токенов'),('precached_prompt_tokens','токены из кеша')]:
                    if key in call.get('tokens',{}):v['details'].append((f'Ответ {i} · {label}',str(call['tokens'][key])))
            if a.get('error'):v['details'].append(('Ошибка',a['error']))
            if kind in ['ocr','barcode']:
                v['tokens']=None;v['summary']='Без расхода лимитов ИИ'
                v['details']=[(k,value) for k,value in v['details'] if k not in ['Источник лимита','Токены GigaChat']]
                v['details'].append(('Токены','Не используются'))
        if v['status']:v['status_label']=STATUSES.get(v['status'],v['status'])
        results.append(v)
    return results


def export(entries,format):
    headings=['ID','Дата и время, МСК','Раздел','Действие','Продукт / объект','Сводка','Статус','Фактические токены GigaChat','Токены: известная часть','Изменения','Подробности','Состояние до (JSON)','Состояние после (JSON)']
    rows=[];changes=[['ID','Поле','До','После']]
    for e in entries:
        tokens=e['tokens'] or {}
        rows.append([e['key'],e['timestamp'],e['category_label'],e['title'],e['name'],e['summary'],e['status_label'],tokens.get('total') if tokens.get('total') is not None else '',tokens.get('known') if tokens.get('partial') else '',
            '\n'.join(f"{c['label']}: {c['before']} → {c['after']}" for c in e['changes']), '\n'.join(f'{k}: {v}' for k,v in e['details']),json.dumps(e['before'],ensure_ascii=False),json.dumps(e['after'],ensure_ascii=False)])
        changes.extend([[e['key'],c['label'],c['before'],c['after']] for c in e['changes']])
    if format=='csv':
        out=io.StringIO();csv.writer(out,delimiter=';',lineterminator='\r\n').writerows([[stock_history.safe_cell(v) for v in row] for row in [headings,*rows]])
        return out.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    book=Workbook();book.remove(book.active)
    for title,data in [('История',[headings,*rows]),('Изменения',changes)]:
        sheet=book.create_sheet(title)
        for row in data:
            sheet.append([str(v) for v in row])
            for cell in sheet[sheet.max_row]:cell.data_type='s'
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for i in range(1,len(data[0])+1):sheet.column_dimensions[get_column_letter(i)].width=32
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='174B3B')
        for row in sheet:
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    out=io.BytesIO();book.save(out)
    return out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
