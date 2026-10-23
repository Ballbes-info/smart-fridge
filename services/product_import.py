"""Bounded, offline document readers. Files never execute code or create products directly."""
import csv
import io
import json
from pathlib import Path
import re
import zipfile
from datetime import date, datetime
from defusedxml import ElementTree as ET
from services import product_service as service

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 100
EXTENSIONS = {'xlsx','xls','csv','tsv','txt','json','docx','pdf','odt','ods','rtf'}
LABELS = dict(name='Название', brand='Марка', manufacturer='Производитель', quantity='Количество', unit='Единица',
    expiry_date='Годен до', manufactured_date='Дата изготовления', opened_date='Дата вскрытия',
    shelf_value='Срок хранения', shelf_unit='Единица срока', price='Цена партии', note='Заметка',
    ingredients='Состав', nutrition_basis='Основание КБЖУ', portion='Размер порции', calories='Калории',
    protein='Белки', fat='Жиры', carbs='Углеводы', other_info='Другие сведения', barcode='Штрихкод',
    category='Категория', storage_id='Место хранения')

def norm(value):
    return re.sub(r'[^\w]+', '', str(value).casefold().replace('ё','е'))

ALIASES = {norm(key):key for key in LABELS}
ALIASES.update({norm(value):key for key,value in LABELS.items()})
for key, aliases in {
    'name':['наименование','товар','продукт','наименование продукта','product','product name'],
    'brand':['бренд','фирма'], 'quantity':['кол-во','вес','масса','объем','количество продукта'],
    'unit':['ед','ед. изм.','единица измерения','units'], 'expiry_date':['срок годности','дата годности','expiry','годендо'],
    'price':['цена','стоимость','цена руб','цена ₽'], 'note':['примечание','комментарий'],
    'calories':['ккал','калорийность','энергетическая ценность ккал'], 'protein':['белок','белки г'],
    'fat':['жир','жиры г'], 'carbs':['углеводы г'], 'storage_id':['место','хранение','storage'],
}.items():
    for alias in aliases: ALIASES[norm(alias)] = key
UNITS = {norm(k):k for k in service.UNITS}
UNITS.update({norm(v[0]):k for k,v in service.UNITS.items()})
UNITS.update({'литр':'l','литра':'l','литров':'l','килограмм':'kg','грамм':'g','штука':'pcs','штук':'pcs','упаковка':'pack'})

class ImportProblem(ValueError): pass

def scalar(value):
    if value is None:return ''
    if isinstance(value, (datetime,date)):return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    if isinstance(value, bool) or isinstance(value,(dict,list)):raise ImportProblem('В полях продукта должны быть текст, число или дата, без вложенных объектов.')
    if isinstance(value,float) and value.is_integer():value=int(value)
    result=str(value).strip()
    if len(result)>10000:raise ImportProblem('Одно из полей слишком длинное: не более 10 000 символов.')
    return result

def normalize_record(raw):
    result={}
    for key,value in raw.items():
        canonical=ALIASES.get(norm(key))
        if canonical:
            if canonical in result:raise ImportProblem(f'Повторяется поле «{LABELS[canonical]}».')
            result[canonical]=scalar(value)
    if not result:raise ImportProblem('Не найдены поля продукта. Используйте названия столбцов из шаблона.')
    for key in ['expiry_date','manufactured_date','opened_date']:
        value=result.get(key,'')
        if re.fullmatch(r'\d{1,2}[./]\d{1,2}[./]\d{4}',value):
            try:result[key]=datetime.strptime(value.replace('/','.'),'%d.%m.%Y').date().isoformat()
            except ValueError:pass
    combined=re.fullmatch(r'(\d+(?:[.,]\d+)?)\s*([\w.]+)',result.get('quantity',''))
    if combined and norm(combined[2]) in UNITS and not result.get('unit'):
        result['quantity']=combined[1];result['unit']=UNITS[norm(combined[2])]
    for key,mapping in [('unit',UNITS),('category',{norm(v):k for k,v in service.CATEGORIES.items()}),
                        ('nutrition_basis',{norm(v):k for k,v in service.BASES.items()}),
                        ('shelf_unit',{norm(v):k for k,v in service.SHELF_UNITS.items()})]:
        if result.get(key):result[key]=mapping.get(norm(result[key]),result[key])
    return result

def decode(data):
    if data.startswith((b'\xff\xfe',b'\xfe\xff')):return data.decode('utf-16')
    try:return data.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:return data.decode('cp1251')
        except UnicodeDecodeError:raise ImportProblem('Не удалось прочитать кодировку текста. Сохраните файл в UTF-8.')

def archive(data):
    z=zipfile.ZipFile(io.BytesIO(data))
    files=z.infolist()
    if len(files)>500 or sum(i.file_size for i in files)>20*1024*1024 or any(i.flag_bits&1 for i in files):
        z.close();raise ImportProblem('Архив документа слишком большой или зашифрован.')
    return z

def table_records(table):
    table=[[scalar(v) for v in row] for row in table if any(v is not None and str(v).strip() for v in row)]
    if not table:return []
    if len(table)>501 or any(len(row)>40 for row in table):raise ImportProblem('Таблица слишком велика: до 500 строк и 40 столбцов.')
    # Vertical form: one named field per row, including a name field.
    if all(len(row)==2 and norm(row[0]) in ALIASES for row in table) and any(ALIASES[norm(r[0])]=='name' for r in table):
        keys=[ALIASES[norm(r[0])] for r in table]
        if len(set(keys))!=len(keys):raise ImportProblem('В вертикальной таблице повторяются поля. Для нескольких продуктов используйте строки и заголовки.')
        return [normalize_record(dict(table))]
    header=table[0];keys=[ALIASES.get(norm(v)) for v in header]
    if 'name' not in keys:raise ImportProblem('В таблице нужен столбец «Название». Скачайте шаблон или проверьте первую строку.')
    recognized=[k for k in keys if k]
    if len(recognized)!=len(set(recognized)):raise ImportProblem('В таблице повторяются столбцы одного поля.')
    output=[]
    for row in table[1:]:
        if len(row)>len(header) and any(row[len(header):]):raise ImportProblem('Количество значений в строке превышает количество заголовков.')
        record={header[i]:value for i,value in enumerate(row[:len(header)]) if keys[i]}
        if any(record.values()):output.append(normalize_record(record))
    return output

def text_records(content):
    if len(content)>200000:raise ImportProblem('Слишком много текста: не более 200 000 символов.')
    content=content.replace('\r\n','\n').replace('\r','\n').strip()
    if not content:return []
    lines=[line.strip() for line in content.splitlines() if line.strip()]
    if any(sep in lines[0] for sep in [';','\t','|',',']):
        try:
            dialect=csv.Sniffer().sniff('\n'.join(lines[:10]),delimiters=';\t|,')
            table=list(csv.reader(io.StringIO(content),dialect))
            if any(ALIASES.get(norm(v))=='name' for v in table[0]):return table_records(table)
        except csv.Error:pass
    output=[];record={}
    for line in lines:
        match=re.match(r'^([^:]{1,50}):\s*(.*)$',line)
        key=ALIASES.get(norm(match[1])) if match else None
        if key:
            if key in record:
                if key=='name':output.append(normalize_record(record));record={}
                else:raise ImportProblem(f'Поле «{LABELS[key]}» повторяется. Начните следующий продукт с «Название:».')
            record[key]=match[2]
        else:
            if record:raise ImportProblem('Не удалось разделить текст на поля. Используйте строки «Название: …», «Количество: …» или шаблон таблицы.')
            line=re.sub(r'^(?:[-•]|\d+[.)])\s+','',line)
            match=re.fullmatch(r'(.+?)\s+(\d+(?:[.,]\d+)?)\s*(кг|г|мл|л|шт\.?|уп\.?)',line,re.I)
            output.append(normalize_record({'name':match[1], 'quantity':match[2], 'unit':match[3]}) if match else {'name':line})
    if record:output.append(normalize_record(record))
    return output

def xml_content(data,ext):
    with archive(data) as z:
        root=ET.fromstring(z.read('word/document.xml' if ext=='docx' else 'content.xml'))
    local=lambda node:node.tag.rsplit('}',1)[-1]
    def words(node):
        if ext=='docx':return ''.join(n.text or '' for n in node.iter() if local(n)=='t')
        if local(node) in ('p','h'):return ''.join(node.itertext())
        return '\n'.join(''.join(n.itertext()) for n in node.iter() if local(n)=='p')
    output=[];paragraphs=[]
    container=next((n for n in root.iter() if local(n)==('body' if ext=='docx' else 'text' if ext=='odt' else 'spreadsheet')),root)
    for node in container:
        kind=local(node)
        if kind in ('tbl','table'):
            if paragraphs:output.extend(text_records('\n'.join(paragraphs)));paragraphs=[]
            rows=[]
            for row in node:
                if local(row) not in ('tr','table-row'):continue
                repeat=int(next((v for k,v in row.attrib.items() if k.endswith('number-rows-repeated')),'1'))
                cells=[]
                for cell in row:
                    if local(cell) not in ('tc','table-cell','covered-table-cell'):continue
                    count=int(next((v for k,v in cell.attrib.items() if k.endswith('number-columns-repeated')),'1'))
                    value=words(cell)
                    if count>40:
                        if value:raise ImportProblem('В таблице слишком много повторяющихся столбцов.')
                        continue
                    cells.extend([value]*count)
                if any(cells):
                    if repeat>100 or len(rows)+repeat>500:raise ImportProblem('В таблице слишком много строк.')
                    rows.extend([cells]*repeat)
            output.extend(table_records(rows))
        elif kind in ('p','h'):
            paragraphs.append(words(node))
    if paragraphs:output.extend(text_records('\n'.join(paragraphs)))
    return output

def read_file(filename,data):
    ext=Path(filename).suffix.lower().lstrip('.')
    if ext not in EXTENSIONS:raise ImportProblem('Этот формат пока не поддерживается. Выберите файл из списка форматов.')
    if not data:raise ImportProblem('Файл пустой.')
    if len(data)>MAX_BYTES:raise ImportProblem('Размер файла — не более 5 МБ.')
    try:
        if ext=='json':
            payload=json.loads(decode(data))
            if isinstance(payload,dict):
                payload=next((payload[k] for k in ['products','продукты','товары'] if k in payload),payload)
            if isinstance(payload,dict):payload=[payload]
            if not isinstance(payload,list) or any(not isinstance(item,dict) for item in payload):raise ImportProblem('JSON должен содержать продукт или список объектов продуктов.')
            if len(payload)>MAX_ROWS:raise ImportProblem('За один импорт можно добавить не более 100 продуктов.')
            rows=[normalize_record(item) for item in payload]
        elif ext in ('xlsx','xls'):
            rows=[]
            if ext=='xlsx':
                from openpyxl import load_workbook
                with archive(data):pass
                book=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
                try:
                    if len(book.worksheets)>5:raise ImportProblem('В книге должно быть не более 5 листов.')
                    for sheet in book:
                        table=[]
                        for index,row in enumerate(sheet.iter_rows()):
                            if index>500 or len(row)>40:raise ImportProblem('Лист слишком большой: до 500 строк и 40 столбцов.')
                            if any(cell.data_type=='f' for cell in row):raise ImportProblem('Формулы не импортируются. Сохраните в таблице значения вместо формул.')
                            table.append([cell.value for cell in row])
                        rows.extend(table_records(table))
                finally:book.close()
            else:
                import xlrd
                book=xlrd.open_workbook(file_contents=data,on_demand=True)
                try:
                    if book.nsheets>5:raise ImportProblem('В книге должно быть не более 5 листов.')
                    for sheet in book.sheets():
                        if sheet.nrows>500 or sheet.ncols>40:raise ImportProblem('Лист слишком большой: до 500 строк и 40 столбцов.')
                        rows.extend(table_records([[xlrd.xldate.xldate_as_datetime(c.value,book.datemode) if c.ctype==xlrd.XL_CELL_DATE else c.value for c in sheet.row(i)] for i in range(sheet.nrows)]))
                finally:book.release_resources()
        elif ext in ('docx','odt','ods'):rows=xml_content(data,ext)
        elif ext=='pdf':
            from pypdf import PdfReader
            book=PdfReader(io.BytesIO(data))
            if book.is_encrypted:raise ImportProblem('Снимите пароль с PDF перед загрузкой.')
            if len(book.pages)>20:raise ImportProblem('В PDF должно быть не более 20 страниц.')
            content='\n'.join(page.extract_text(extraction_mode='layout') or '' for page in book.pages)
            # Layout extraction keeps columns separated without guessing single spaces in names.
            lines=content.splitlines()
            if any(ALIASES.get(norm(part))=='name' for line in lines[:5] for part in re.split(r'\s{2,}',line.strip())):
                content='\n'.join('\t'.join(re.split(r' {2,}',line.strip())) for line in lines if line.strip())
            if not content.strip():raise ImportProblem('В PDF нет текстового слоя. Загрузите документ с текстом или добавьте фото через OCR / ИИ.')
            rows=text_records(content)
        elif ext=='rtf':
            from striprtf.striprtf import rtf_to_text
            rows=text_records(rtf_to_text(decode(data)))
        else:
            content=decode(data)
            if ext in ('csv','tsv'):
                try:dialect=csv.Sniffer().sniff(content[:10000],delimiters=';\t,|')
                except csv.Error:raise ImportProblem('Не удалось определить разделитель. Используйте CSV-шаблон.')
                rows=table_records(list(csv.reader(io.StringIO(content),dialect)))
            else:rows=text_records(content)
    except ImportProblem:raise
    except Exception as exc:
        raise ImportProblem('Не удалось прочитать файл. Проверьте формат, снимите пароль или сохраните документ заново.') from exc
    if not rows:raise ImportProblem('В файле не найдено продуктов.')
    if len(rows)>MAX_ROWS:raise ImportProblem('За один импорт можно добавить не более 100 продуктов. Разделите файл.')
    return rows
